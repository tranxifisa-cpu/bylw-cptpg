#!/usr/bin/env python
"""Minimal, safe edit: fix reference list only. Never touches inline text."""
from docx import Document
from docx.shared import Pt
from copy import deepcopy
import re

SRC = r'd:\project\bylw\output_paper\doc\董祉含_开题报告书和申请表_带引用编号.docx'
DST = r'd:\project\bylw\output_paper\doc\董祉含_开题报告书和申请表_已修复.docx'

doc = Document(SRC)

# ============================================================
# Step 1: Identify reference paragraphs by number
# ============================================================
ref_paras = {}  # number -> (paragraph_element, parent)
for para in doc.paragraphs:
    text = ''.join(r.text for r in para.runs if r.text).strip()
    m = re.match(r'^\[(\d+)\]', text)
    if m:
        ref_paras[int(m.group(1))] = para

print(f'Found {len(ref_paras)} reference entries: {sorted(ref_paras.keys())}')

# ============================================================
# Step 2: Delete unused references
# ============================================================
# These are NOT cited in the current proposal text:
UNUSED = [1, 2, 8, 9, 10, 11, 17, 18, 19]

for num in UNUSED:
    if num in ref_paras:
        p = ref_paras[num]._element
        parent = p.getparent()
        parent.remove(p)
        print(f'  Deleted unused [{num}]')

# ============================================================
# Step 3: Add Arkes 2010 after Arkes 2008 (which is [5])
# ============================================================
arkes_2008 = ref_paras.get(5)
if arkes_2008:
    # Create a new paragraph by deep-copying Arkes 2008
    new_p = deepcopy(arkes_2008)
    # Clear text from all runs
    for run in new_p.runs:
        run.text = ''
    # Set new text in first run
    text_2010 = (
        '[5b] Arkes, H. R., Hirshleifer, D., Jiang, D., & Lim, S. S. (2010). '
        'A cross-cultural study of reference point adaptation: Evidence from China, Korea, and the US. '
        'Organizational Behavior and Human Decision Processes, 112(2), 99–111.'
    )
    new_p.runs[0].text = text_2010
    new_p.runs[0].font.name = 'Times New Roman'
    new_p.runs[0].font.size = Pt(10)
    # Remove extra runs
    for run in new_p.runs[1:]:
        run.text = ''

    # Insert after Arkes 2008
    arkes_2008._element.addnext(new_p._element)
    print('  Added Arkes 2010 as [5b] after [5]')

# ============================================================
# Step 4: Renumber the reference list to be sequential
# ============================================================
# Current numbers after cleanup: 3,4,5,5b,6,7,12,13,14,15,16,20,21
# Renumber to 1-13 sequentially
RENUMBER = {
    3: 1, 4: 2, 5: 3, '5b': 4, 6: 5, 7: 6,
    12: 7, 13: 8, 14: 9, 15: 10, 16: 11, 20: 12, 21: 13
}

# Re-read paragraphs after modifications
for para in doc.paragraphs:
    text = ''.join(r.text for r in para.runs if r.text).strip()
    m = re.match(r'^\[(\d+[a-z]?)\]', text)
    if not m:
        continue
    key_str = m.group(1)
    key = int(key_str) if key_str.isdigit() else key_str
    if key in RENUMBER:
        new_num = RENUMBER[key]
        # Replace the reference number in each run's text
        for run in para.runs:
            if run.text and run.text.startswith(f'[{key_str}]'):
                run.text = run.text.replace(f'[{key_str}]', f'[{new_num}]', 1)
                break

print(f'  Renumbered references to [1]-[{max(RENUMBER.values())}]')

# ============================================================
# Save
# ============================================================
doc.save(DST)
print(f'\nSaved: {DST}')

# ============================================================
# Report inline citations that need manual addition
# ============================================================
print('\n' + '='*60)
print('INLINE CITATIONS TO ADD MANUALLY (编号对应新参考文献):')
print('='*60)
print("""
在正文中找到以下位置，手动添加引用编号（不要改动我的脚本，直接在Word里加）：

[1] Kahneman & Tversky (1979)  → 在 P69 "Kahneman和Tversky（1992）" 处改为 "Kahneman和Tversky[1][2]"
    注：原文(P69)引的是1992版CPT，需确认是否同时引了1979版PT
[2] Tversky & Kahneman (1992)  → 同上，CPT的原典

[3] Arkes et al. (2008)        → P71 "Arkes等（2008）" 后加 [3]
                                  P95 "Arkes、Hirshleifer、Jiang和Lim（2008）" 后加 [3]

[4] Arkes et al. (2010)        → P71 "跨文化样本...重复验证" 后加 [4]（如该段提到了跨文化）

[5] He & Yang (2019)           → P72 "He和Yang（2019）" 后加 [5]
                                  P96 "He和Yang（2019）" 后加 [5]

[6] Qin et al. (2022)          → P72 "Qin等（2022）" 后加 [6]
                                  P97 "Qin、Kong和Yue（2022）" 后加 [6]

[7] Prashanth et al. (2016)    → P70 "Prashanth等（2016）" 后加 [7]
                                  P71 "Prashanth等（2016）" 后加 [7]
                                  P91 "Prashanth等" 后加 [7]

[8] Lepel & Barakat (2026)     → P70 "Lepel和Barakat（2026）" 后加 [8]
                                  P71 "Lepel和Barakat（2026）" 后加 [8]
                                  P91 "Lepel和Barakat" 后加 [8]

[9] Fei et al. (2020)          → P100 "Fei、Yang、Wang和Xie（2020）" 后加 [9]
                                  P102 "Fei等（2020）" 后加 [9]

[10] Ghadimi & Lan (2013)      → P102 "Ghadimi和Lan（2013）" 后加 [10]

[11] Andre & Coqueret (2021)   → P101 "André和Coqueret（2021）" 后加 [11]
                                  P102 "André和Coqueret（2021）" 后加 [11]

[12] Ramasubramanian et al. (2021) → P70 "Ramasubramanian等（2021）" 后加 [12]
                                      P92 "Ramasubramanian" 后加 [12]

[13] Lalmohammed et al. (2025) → P70 "Lalmohammed等（2025）" 后加 [13]
                                  P92 "Lalmohammed等" 后加 [13]
""")
