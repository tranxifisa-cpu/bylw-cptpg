#!/usr/bin/env python
"""Add in-text citation numbers to proposal DOCX. Preserves all formatting.
Reads the original DOCX, finds author-year mentions, inserts [N] markers.
Also cleans up the reference list: removes unused, adds missing."""

from docx import Document
from docx.shared import Pt
from copy import deepcopy
import re

SRC = r'd:\project\bylw\output_paper\doc\董祉含_开题报告书和申请表.docx'
DST = r'd:\project\bylw\output_paper\doc\董祉含_开题报告书和申请表_带引用编号.docx'

doc = Document(SRC)

# ============================================================
# Step 1: Map citation patterns to reference numbers
# ============================================================
# Reference list mapping (existing numbers -> what they are):
# [3] Kahneman & Tversky 1979
# [4] Tversky & Kahneman 1992
# [5] Arkes et al. 2008
# [12] Prashanth et al. 2016
# [13] Lepel & Barakat 2026
# [20] Ramasubramanian et al. 2021
# [21] Lalmohammed et al. 2025
# [6] He & Yang 2019
# [7] Qin et al. 2022
# [14] Fei et al. 2020
# [15] Ghadimi & Lan 2013
# [16] Andre & Coqueret 2021
#
# To add: Arkes et al. 2010
# To delete: [1][2][8][9][10][11][17][18][19]

# Pattern -> citation number mapping (order of first appearance in text)
PATTERN_MAP = [
    # (regex_pattern, citation_number)
    (r'Kahneman.*?Tversky.*?1992', '[3][4]'),  # combined CPT ref
    (r'(?<!Lev )Tversky.*?Kahneman.*?1992', '[4]'),  # CPT
    (r'Kahneman.*?Tversky.*?1979', '[3]'),  # PT original
    (r'Prashanth.*?2016', '[12]'),
    (r'Lepel.*?Barakat.*?2026', '[13]'),
    (r'Ramasubramanian.*?2021', '[20]'),
    (r'Lalmohammed.*?2025', '[21]'),
    (r'Arkes.*?2010', '[新增]'),  # to be added
    (r'Arkes.*?2008', '[5]'),
    (r'He.*?Yang.*?2019', '[6]'),
    (r'Qin.*?2022', '[7]'),
    (r'Fei.*?2020', '[14]'),
    (r'Ghadimi.*?Lan.*?2013', '[15]'),
    (r'Andr.*?Coqueret.*?2021', '[16]'),
]

# ============================================================
# Step 2: Add citation numbers in text paragraphs
# ============================================================
def get_para_full_text(para):
    """Get full text of a paragraph."""
    return ''.join(run.text for run in para.runs if run.text)

def add_citation_to_para(para):
    """Find citation patterns in a paragraph and add [N] markers."""
    full_text = get_para_full_text(para)
    if not full_text.strip():
        return False

    modified = False

    for pattern, cit_num in PATTERN_MAP:
        if cit_num == '[新增]':
            continue  # skip, handle separately

        match = re.search(pattern, full_text)
        if not match:
            continue

        # Found a citation. Find the exact position in the paragraph runs.
        target = match.group()
        target_end = match.end()

        # Find which run contains the end of the match
        char_count = 0
        for run in para.runs:
            if not run.text:
                char_count += 0
                continue
            run_start = char_count
            run_end = char_count + len(run.text)
            char_count = run_end

            if target_end <= run_end and target_end > run_start:
                # The citation ends within this run
                local_pos = target_end - run_start
                # Insert [N] right after the citation
                old_text = run.text
                run.text = old_text[:local_pos] + cit_num + old_text[local_pos:]
                modified = True
                break

    return modified

# Process all paragraphs
count = 0
for para in doc.paragraphs:
    if add_citation_to_para(para):
        count += 1

print(f'Added citations to {count} paragraphs.')

# ============================================================
# Step 3: Fix the reference list
# ============================================================
# Find the reference section (paragraphs starting with [N])
# Delete unused: [1][2][8][9][10][11][17][18][19]
# Add: Arkes et al. 2010 as new entry (renumber appropriately)

ref_entries = {}  # number -> paragraph
unused = {1, 2, 8, 9, 10, 11, 17, 18, 19}

for para in doc.paragraphs:
    text = get_para_full_text(para).strip()
    m = re.match(r'^\[(\d+)\]', text)
    if m:
        num = int(m.group(1))
        ref_entries[num] = para

# Delete unused references
for num in sorted(unused, reverse=True):
    if num in ref_entries:
        p = ref_entries[num]._element
        p.getparent().remove(p)
        print(f'Removed unused reference [{num}]')

# Add Arkes 2010 - insert after [5] (Arkes 2008)
# Find the Arkes 2008 paragraph and insert after it
arkes_2008_para = ref_entries.get(5)
if arkes_2008_para:
    new_p = deepcopy(arkes_2008_para)
    # Clear all runs and add new text
    for run in new_p.runs:
        run.text = ''
    new_p.runs[0].text = '[5-2] Arkes, H. R., Hirshleifer, D., Jiang, D., & Lim, S. S. (2010). A cross-cultural study of reference point adaptation: Evidence from China, Korea, and the US. Organizational Behavior and Human Decision Processes, 112(2), 99–111.'
    # Set font
    for run in new_p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10.5) if run.font.size is None else run.font.size

    arkes_2008_para._element.addnext(new_p._element)
    print('Added Arkes 2010 after Arkes 2008.')

# ============================================================
# Step 4: Add "Arkes et al. (2010)" citation in the text
# ============================================================
arkes_2010_added = False
for para in doc.paragraphs:
    full = get_para_full_text(para)
    if 'Arkes等（2008）' in full and '跨文化' in full and not arkes_2010_added:
        # This paragraph discusses cross-cultural replication
        # Find and modify to add Arkes 2010 citation
        for run in para.runs:
            if '重复验证' in (run.text or '') or '跨文化' in (run.text or ''):
                run.text = run.text.replace('重复验证', '重复验证[新增Arkes2010]')
                arkes_2010_added = True
                break
        if arkes_2010_added:
            break

# ============================================================
# Save
# ============================================================
doc.save(DST)
print(f'\nSaved: {DST}')
print('Done. Please verify citation numbers manually.')
